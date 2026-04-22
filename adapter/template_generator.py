"""
adapter/template_generator.py

Generate a ready-to-fill Excel template (.xlsx) from a mapping.yml configuration.

The generated workbook contains:
  - A header row with one column per mapped field (required fields highlighted).
  - An example row with realistic placeholder values.
  - Data validation dropdowns on the Status and Party columns.
  - Frozen first row and auto-sized column widths.
  - A "README" sheet with field descriptions.

Requires openpyxl:  pip install openpyxl
"""

from __future__ import annotations

try:
    import openpyxl
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side, numbers
    )
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# ── Field metadata ────────────────────────────────────────

# For each canonical runbook field: (required, description, example_value)
FIELD_META: dict[str, tuple[bool, str, str]] = {
    "task":         (True,  "Task description (what needs to be done)",          "Deploy application to PROD"),
    "status":       (True,  "Current task status",                                "Not Started"),
    "startTime":    (False, "Scheduled start time (YYYY-MM-DDTHH:MM or HH:MM)",  "2026-04-15T09:00"),
    "endTime":      (False, "Scheduled end time   (YYYY-MM-DDTHH:MM or HH:MM)",  "2026-04-15T09:30"),
    "item":         (False, "Short reference/ticket ID",                          "TASK-001"),
    "assignee":     (False, "Person or team responsible",                         "John Smith"),
    "taskId":       (False, "Numeric task identifier (used for ordering/linking)","42"),
    "system":       (False, "Impacted system or application",                     "MX.3 / PROD"),
    "party":        (False, "Responsible party (Client / Murex / Joint)",         "Client"),
    "estimatedEnd": (False, "Originally planned end time (frozen at import)",     "2026-04-15T09:25"),
    "comment":      (False, "Free-text notes visible in the dashboard",           "Check with ops team first"),
}

# Canonical status labels that are always included in the dropdown
CANONICAL_STATUSES = [
    "Not Started", "In Progress", "Completed", "Blocking", "Unneeded",
]

PARTY_OPTIONS = ["Client", "Murex", "Joint"]

# Colours — only defined when openpyxl is available
if HAS_OPENPYXL:
    HEADER_REQUIRED_FILL = PatternFill("solid", fgColor="FFF9C4")   # yellow
    HEADER_OPTIONAL_FILL = PatternFill("solid", fgColor="E8F5E9")   # light green
    HEADER_FONT          = Font(bold=True, size=11)
    EXAMPLE_FONT         = Font(italic=True, color="555555", size=10)
    README_HEADER_FONT   = Font(bold=True, size=12)
    THIN_BORDER          = Border(
        bottom=Side(style="thin", color="AAAAAA"),
        right=Side(style="thin",  color="CCCCCC"),
    )
else:
    HEADER_REQUIRED_FILL = HEADER_OPTIONAL_FILL = HEADER_FONT = None
    EXAMPLE_FONT = README_HEADER_FONT = THIN_BORDER = None


# ── Internal helpers ──────────────────────────────────────

def _status_dropdown_formula(mapping: dict) -> str:
    """Build the quoted comma-list for the Status column dropdown."""
    src_values: list[str] = []
    status_map = mapping.get("status_mapping", {})
    if status_map:
        src_values = list(status_map.keys())
    if not src_values:
        src_values = CANONICAL_STATUSES
    # Excel data-validation formula: "v1,v2,v3" (max ~255 chars)
    joined = ",".join(src_values[:20])  # cap to avoid Excel limit
    return f'"{joined}"'


def _col_width(text: str) -> float:
    """Estimate a comfortable column width from sample text length."""
    return min(max(len(text) + 4, 12), 40) * 1.1


# ── Public API ────────────────────────────────────────────

def generate_template(mapping: dict, output_path: str) -> None:
    """
    Write an Excel template to *output_path* based on *mapping*.

    :param mapping: dict loaded from mapping.yml / mapping.json
    :param output_path: destination .xlsx file path
    :raises ImportError: if openpyxl is not installed
    """
    if not HAS_OPENPYXL:
        raise ImportError(
            "openpyxl is required to generate Excel templates.\n"
            "Install with:  pip install openpyxl"
        )

    cols_cfg     = mapping.get("columns", {})
    cat_col      = mapping.get("category_column")
    default_cat  = mapping.get("default_category", "Tasks")
    status_map   = mapping.get("status_mapping", {})

    # Build ordered list of output columns
    # category column goes first (if defined), then mapped fields in FIELD_META order
    output_cols: list[tuple[str, str]] = []   # (field_key, header_label)

    if cat_col:
        output_cols.append(("__category__", cat_col))

    for field in FIELD_META:
        col_label = cols_cfg.get(field)
        if col_label:
            output_cols.append((field, col_label))

    if not output_cols:
        raise ValueError(
            "No columns found in mapping. "
            "Make sure mapping.yml has a 'columns:' section."
        )

    wb = openpyxl.Workbook()

    # ── Main "Runbook" sheet ──────────────────────────────
    ws = wb.active
    ws.title = "Runbook"

    # Header row
    for col_idx, (field, label) in enumerate(output_cols, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        required = field != "__category__" and FIELD_META.get(field, (False,))[0]
        cell.fill  = HEADER_REQUIRED_FILL if required else HEADER_OPTIONAL_FILL
        cell.font  = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
        ws.row_dimensions[1].height = 30

    # Example row
    example_row: list[str] = []
    for field, label in output_cols:
        if field == "__category__":
            example_row.append(default_cat)
        else:
            meta = FIELD_META.get(field)
            if field == "status" and status_map:
                # Use first source-side status label from the mapping
                example_row.append(next(iter(status_map.keys()), "Not Started"))
            else:
                example_row.append(meta[2] if meta else "")

    for col_idx, value in enumerate(example_row, start=1):
        cell = ws.cell(row=2, column=col_idx, value=value)
        cell.font   = EXAMPLE_FONT
        cell.border = THIN_BORDER

    # Data validation — Status column
    status_col_idx = next(
        (i + 1 for i, (f, _) in enumerate(output_cols) if f == "status"), None
    )
    if status_col_idx:
        col_letter = get_column_letter(status_col_idx)
        dv_status = DataValidation(
            type="list",
            formula1=_status_dropdown_formula(mapping),
            allow_blank=True,
            showDropDown=False,
        )
        dv_status.sqref = f"{col_letter}2:{col_letter}1000"
        ws.add_data_validation(dv_status)

    # Data validation — Party column
    party_col_idx = next(
        (i + 1 for i, (f, _) in enumerate(output_cols) if f == "party"), None
    )
    if party_col_idx:
        col_letter = get_column_letter(party_col_idx)
        dv_party = DataValidation(
            type="list",
            formula1='"' + ",".join(PARTY_OPTIONS) + '"',
            allow_blank=True,
            showDropDown=False,
        )
        dv_party.sqref = f"{col_letter}2:{col_letter}1000"
        ws.add_data_validation(dv_party)

    # Auto-size columns
    for col_idx, (field, label) in enumerate(output_cols, start=1):
        meta  = FIELD_META.get(field)
        sample = meta[2] if meta else label
        width  = _col_width(max(label, sample, key=len))
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze header row
    ws.freeze_panes = "A2"

    # ── README sheet ──────────────────────────────────────
    ws_readme = wb.create_sheet("README")
    ws_readme.column_dimensions["A"].width = 18
    ws_readme.column_dimensions["B"].width = 16
    ws_readme.column_dimensions["C"].width = 50
    ws_readme.column_dimensions["D"].width = 32

    readme_headers = ["Field", "Required", "Description", "Example"]
    for col_idx, h in enumerate(readme_headers, start=1):
        cell = ws_readme.cell(row=1, column=col_idx, value=h)
        cell.font  = README_HEADER_FONT
        cell.fill  = PatternFill("solid", fgColor="E3F2FD")
        cell.border = THIN_BORDER

    readme_row = 2
    if cat_col:
        ws_readme.cell(row=readme_row, column=1, value=cat_col)
        ws_readme.cell(row=readme_row, column=2, value="Yes")
        ws_readme.cell(row=readme_row, column=3, value="Groups tasks into collapsible sections in the dashboard")
        ws_readme.cell(row=readme_row, column=4, value=default_cat)
        readme_row += 1

    for field, label in output_cols:
        if field == "__category__":
            continue
        meta = FIELD_META.get(field)
        if not meta:
            continue
        required, desc, example = meta
        ws_readme.cell(row=readme_row, column=1, value=label)
        ws_readme.cell(row=readme_row, column=2, value="Yes" if required else "No")
        ws_readme.cell(row=readme_row, column=3, value=desc)
        ws_readme.cell(row=readme_row, column=4, value=example)
        readme_row += 1

    ws_readme.freeze_panes = "A2"

    wb.save(output_path)
